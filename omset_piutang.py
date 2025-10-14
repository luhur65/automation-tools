import re
import os
from dotenv import load_dotenv
from datetime import datetime, timedelta
from playwright.sync_api import sync_playwright, Page
import openpyxl
from openpyxl.formatting.rule import Rule

load_dotenv() # Muat variabel dari file .env

USERNAME = os.getenv("APP_USERNAME")
PASSWORD = os.getenv("APP_PASSWORD")
URL = os.getenv("APP_TARGET_URL")
EXCEL_FILE = "cek_harian.xlsx"
CABANG = {
    "Medan": "Mdn",
    "Jakarta": "Jkt",
    "Surabaya": "Sby",
    "Makassar": "Mks",
    "Semarang": "Smg"
}

def apply_dynamic_format(cell, value: float):
    """
    Menerapkan format angka dinamis ke sebuah sel.
    Jika nilainya bilangan bulat, formatnya '#,##0'.
    Jika desimal, formatnya '#,##0.##'.
    """
    if value == int(value):
        cell.number_format = '#,##0'
    else:
        cell.number_format = '#,##0.##'


def autosize_columns(ws):
    # set reasonable widths
    widths = {
        "A": 15, "B": 35, "C": 18, "D": 18, "E": 35, "F": 18, "G": 18
    }
    for col, w in widths.items():
        ws.column_dimensions[col].width = w

# === SCRAP OMSET + PROFIT ===
def scrap_omset_profit(page: Page, cabang: str, kode: str) -> list:
    """Mengambil data omset, profit, dan last update untuk satu cabang."""
    page.get_by_role("link", name=f"Omset {cabang}").click()
    page.wait_for_timeout(800)

    omset = page.locator(f'tr.footrow-ltr.ui-widget-content.myfootrowJKT td[aria-describedby="toolbar{kode}_FOmset"]').last.inner_text().strip()
    profit = page.locator(f'tr.footrow-ltr.ui-widget-content.myfootrowJKT td[aria-describedby="toolbar{kode}_FProfit"]').last.inner_text().strip()

    # Cari tanggal last update dari blok tabs
    last_update = "-"
    try:
        # ambil parent container tabs
        tabs_container = page.locator(f"#tabs-{list(CABANG.keys()).index(cabang)+1}")
        html = tabs_container.inner_html()

        match = re.search(
            r"Tanggal Last Update\s*:\s*([\d-]+\s[\d:]+)",
            html
        )
        if match:
            last_update = match.group(1)
    except Exception as e:
        print(f"⚠️ Gagal ambil last update {cabang}: {e}")
        # Jika gagal, biarkan nilainya "-"

    return [cabang, last_update, omset, profit]

# === SCRAP PIUTANG ===
def scrap_piutang(page: Page, cabang: str, kode: str) -> str:
    page.get_by_role("link", name=f"Piutang EMKL {cabang}").click()
    page.wait_for_timeout(800)

    idx = list(CABANG.keys()).index(cabang) + 1
    container_html = page.inner_html(f"div#tabs-{idx}")

    last_update = "-"
    m = re.search(r"Last Update.*?(\d{2}-\d{2}-\d{4}\s\d{2}:\d{2}:\d{2})", container_html, re.S)
    if m:
        last_update = m.group(1)

    return last_update


def run():
    """Fungsi utama untuk menjalankan proses scraping dan penulisan ke Excel."""
    today = datetime.now().strftime("%d-%m-%Y")
    yesterday = (datetime.now() - timedelta(days=1)).strftime("%d-%m-%Y")

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context()
        page = context.new_page()

        # Login
        page.goto(f"{URL}/sys/")
        page.get_by_role("textbox", name="Masukkan userid.....").fill(USERNAME)
        page.get_by_role("textbox", name="Masukkan password.....").fill(PASSWORD)
        page.get_by_role("button", name="Login").click()
        page.wait_for_load_state("networkidle")

        # Masuk ke Omset
        # page.get_by_text("Laporan EMKL", exact=True).click()
        # page.get_by_role("link", name="Laporan Omset", exact=True).click()
        # langsung ke halaman omset
        page.goto(f"{URL}/sys/Omset/index")

        data = []
        for cabang, kode in CABANG.items():
            data.append(scrap_omset_profit(page, cabang, kode))

        # Masuk ke Piutang
        # page.get_by_text("Laporan EMKL", exact=True).click()
        # page.get_by_role("link", name="Laporan Piutang EMKL").click()
        page.goto(f"{URL}/sys/piutangemkl/index")

        piutang_data = {}
        for cabang in CABANG.keys():
            piutang_data[cabang] = scrap_piutang(page, cabang, kode)
        browser.close()

    # === BUAT / BUKA EXCEL ===
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
    except FileNotFoundError:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

    # Buat sheet baru untuk hari ini
    ws = wb.create_sheet(title=today)

    # Header
    ws.append(["Cabang", "Last Update Omset", "Omset", "Profit",
               "Last Update Piutang", "SELISIH OMSET", "SELISIH PROFIT"])

    # Isi data cabang
    for i, row in enumerate(data, start=2):
        cabang, last_update, omset, profit = row

        ws[f"A{i}"] = cabang
        ws[f"B{i}"] = f"Tanggal Last Update : {last_update}"

        # Omset
        omset_val = float(omset.replace(",", "")) if omset else 0
        cell_omset = ws[f"C{i}"]
        cell_omset.value = omset_val
        apply_dynamic_format(cell_omset, omset_val)

        # Profit
        profit_val = float(profit.replace(",", "")) if profit else 0
        cell_profit = ws[f"D{i}"]
        cell_profit.value = profit_val
        apply_dynamic_format(cell_profit, profit_val)

        # ws[f"E{i}"] = piutang_data.get(cabang, "-")  # nanti diisi piutang
        ws[f"E{i}"] = f"Last Update : {piutang_data.get(cabang, '-')}"

        # Formula selisih (kalau sheet kemarin ada)
        if yesterday in wb.sheetnames:
            # SELISIH OMSET
            cell_selisih_omset = ws[f"F{i}"]
            cell_selisih_omset.value = f"=C{i}-'{yesterday}'!C{i}"
            cell_selisih_omset.number_format = '#,##0.##' # Format default
            
            # Buat aturan conditional formatting
            rule_omset = Rule(type="expression", dxfId=0, stopIfTrue=True)
            rule_omset.formula = [f"MOD(F{i},1)=0"] # Cek jika bilangan bulat
            rule_omset.dxf.number_format.format_code = '#,##0' # Format jika benar
            ws.conditional_formatting.add(f"F{i}", rule_omset)

            # SELISIH PROFIT
            cell_selisih_profit = ws[f"G{i}"]
            cell_selisih_profit.value = f"=D{i}-'{yesterday}'!D{i}"
            cell_selisih_profit.number_format = '#,##0.##' # Format default

            # Buat aturan conditional formatting
            rule_profit = Rule(type="expression", dxfId=1, stopIfTrue=True) # dxfId harus beda
            rule_profit.formula = [f"MOD(G{i},1)=0"] # Cek jika bilangan bulat
            rule_profit.dxf.number_format.format_code = '#,##0' # Format jika benar
            ws.conditional_formatting.add(f"G{i}", rule_profit)

    autosize_columns(ws)
    wb.save(EXCEL_FILE)
    print(f"✅ Laporan {today} berhasil disimpan di {EXCEL_FILE}")

if __name__ == "__main__":
    run()