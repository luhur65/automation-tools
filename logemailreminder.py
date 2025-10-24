import os
import re
from datetime import datetime
from dotenv import load_dotenv
from playwright.sync_api import sync_playwright, Page
import openpyxl

load_dotenv()

USERNAME = os.getenv("TRUCKING_USERNAME")
PASSWORD = os.getenv("TRUCKING_PASSWORD")
BASE_URLS = {
    "Medan": os.getenv("URL_MEDAN"),
    "Jakarta": os.getenv("URL_JAKT"),
    "Surabaya": os.getenv("URL_SBY"),
    "Makassar": os.getenv("URL_MKS"),
}
EXCEL_FILE = "cek_log_email_reminder.xlsx"

def cek_log_email(page: Page, cabang: str, base_url: str) -> str:
    # setelah login sudah, langsung ke halaman log reminder email
    url = f"{base_url}/trucking/logreminderemail/index"
    page.goto(url)
    page.wait_for_load_state("networkidle")

    html = page.inner_html("#gview_jqGrid")  # asumsi grid punya id ini
    today_str = datetime.now().strftime("%d-%m-%Y")
    # Cari tanggal hari ini + waktu
    match = re.search(rf"{today_str}\s+\d{{2}}:\d{{2}}:\d{{2}}", html)
    if match:
        return match.group(0)
    return "-"

def login(page: Page, base_url: str):
    page.goto(f"{base_url}/trucking/login")
    page.get_by_role("textbox", name="User ID").fill(USERNAME)
    page.get_by_role("textbox", name="Password").fill(PASSWORD)
    page.get_by_role("button", name="Sign In").click()
    page.wait_for_load_state("networkidle")

def run():
    today = datetime.now().strftime("%d-%m-%Y")
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        results = {}
        for cabang, base_url in BASE_URLS.items():
            context = browser.new_context()
            page = context.new_page()
            try:
                login(page, base_url)
                waktu = cek_log_email(page, cabang, base_url)
                results[cabang] = waktu
            except Exception as e:
                results[cabang] = f"Error: {e}"
            finally:
                context.close()
        browser.close()

    # === Tulis ke Excel ===
    try:
        wb = openpyxl.load_workbook(EXCEL_FILE)
    except FileNotFoundError:
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

    ws = wb.create_sheet(title=today)
    ws.append(["Cabang", "Waktu Pengiriman Email Reminder"])

    for i, (cabang, waktu) in enumerate(results.items(), start=2):
        ws[f"A{i}"] = cabang
        ws[f"B{i}"] = f"Tanggal Kirim : {waktu}"

    wb.save(EXCEL_FILE)
    print(f"✅ Hasil cek log email reminder berhasil disimpan di {EXCEL_FILE}")

if __name__ == "__main__":
    run()
